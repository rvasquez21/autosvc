#!/usr/bin/env python3
"""
Script de diagnóstico para archivos Excel problemáticos
"""

import sys
import os
import pandas as pd
import openpyxl
from datetime import datetime

# Agregar el directorio actual al path para importar app
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

def diagnose_excel_file(file_path):
    """Diagnosticar un archivo Excel específico"""
    
    print(f"🔍 Diagnóstico de archivo: {file_path}")
    print("=" * 60)
    
    if not os.path.exists(file_path):
        print("❌ El archivo no existe")
        return False
    
    file_size = os.path.getsize(file_path)
    print(f"📊 Tamaño del archivo: {file_size} bytes")
    
    if file_size == 0:
        print("❌ El archivo está vacío")
        return False
    
    # Verificar extensión
    file_ext = os.path.splitext(file_path)[1].lower()
    print(f"📁 Extensión: {file_ext}")
    
    if file_ext not in ['.xlsx', '.xls']:
        print("❌ No es un archivo Excel válido")
        return False
    
    # Método 1: Verificar con openpyxl
    print("\n🔧 Método 1: Verificación con openpyxl")
    try:
        workbook = openpyxl.load_workbook(file_path)
        print(f"✅ Archivo válido con openpyxl")
        print(f"📋 Hojas disponibles: {workbook.sheetnames}")
        
        # Información de la primera hoja
        sheet = workbook.active
        print(f"📊 Hoja activa: {sheet.title}")
        print(f"📊 Dimensiones: {sheet.max_row} filas x {sheet.max_column} columnas")
        
        # Mostrar primeras filas
        print("\n📋 Primeras 5 filas:")
        for row_num in range(1, min(6, sheet.max_row + 1)):
            row_data = []
            for col_num in range(1, min(6, sheet.max_column + 1)):
                cell_value = sheet.cell(row=row_num, column=col_num).value
                row_data.append(str(cell_value)[:20] if cell_value else "")
            print(f"   Fila {row_num}: {row_data}")
        
    except Exception as e:
        print(f"❌ Error con openpyxl: {e}")
    
    # Método 2: Verificar con pandas
    print("\n🔧 Método 2: Verificación con pandas")
    try:
        df = pd.read_excel(file_path, engine='openpyxl')
        print(f"✅ Archivo leído con pandas")
        print(f"📊 Dimensiones: {len(df)} filas x {len(df.columns)} columnas")
        print(f"📋 Columnas: {list(df.columns)}")
        
        # Mostrar tipos de datos
        print(f"📊 Tipos de datos:")
        for col in df.columns:
            print(f"   {col}: {df[col].dtype}")
        
        # Mostrar primeras filas
        print(f"\n📋 Primeras 3 filas:")
        print(df.head(3).to_string())
        
    except Exception as e:
        print(f"❌ Error con pandas: {e}")
    
    # Método 3: Verificar como HTML (para archivos .xls antiguos)
    if file_ext == '.xls':
        print("\n🔧 Método 3: Verificación como HTML")
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
                if '<html' in content.lower() or '<table' in content.lower():
                    print("✅ Archivo parece ser HTML")
                    print(f"📊 Contiene HTML: {'<html' in content.lower()}")
                    print(f"📊 Contiene tablas: {'<table' in content.lower()}")
                else:
                    print("❌ No parece ser HTML")
        except Exception as e:
            print(f"❌ Error leyendo como texto: {e}")
    
    return True

def main():
    """Función principal"""
    print("🔧 Diagnóstico de Archivos Excel - AutoSvc")
    print("=" * 60)
    
    # Verificar archivos en uploads
    uploads_dir = "uploads"
    if not os.path.exists(uploads_dir):
        print("❌ Directorio uploads no existe")
        return
    
    files = os.listdir(uploads_dir)
    excel_files = [f for f in files if f.endswith(('.xlsx', '.xls'))]
    
    if not excel_files:
        print("❌ No se encontraron archivos Excel en uploads/")
        return
    
    print(f"📁 Archivos Excel encontrados: {excel_files}")
    print()
    
    for file_name in excel_files:
        file_path = os.path.join(uploads_dir, file_name)
        diagnose_excel_file(file_path)
        print("\n" + "="*60 + "\n")

if __name__ == "__main__":
    main()
